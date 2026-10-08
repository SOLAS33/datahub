"""Fuel sources: pump and heating-oil prices, Irish fuel taxes, crude oil and exchange rates, SEAI heating-fuel and energy-use files, NORA monthly volumes.

Used by Fuelwatch Ireland. Every file is a tidy CSV or JSON built from an official publication, with the publisher's own values (no estimates).
Parsers take plain rows or text so tests can feed them small samples; the collectors only fetch, check and publish.

* `eu_oil_bulletin`: the European Commission Weekly Oil Bulletin history (2005 on): prices with and without taxes for every member state, and Ireland's tax tables.
* `revenue_mot`: Revenue's Mineral Oil Tax rates since November 2008 (from the Excise Duty Rates manual) and the solid fuel and natural gas carbon tax rates.
* `eia_brent`: US EIA Europe Brent spot price, dollars per barrel, daily.
* `ecb_fx`: ECB euro reference rates (US dollar and sterling).
* `seai_fuel`: SEAI domestic fuel price archive, final energy consumption, energy by fuel and the National Energy Balance.
* `nora_volumes`: NORA monthly volumes of oil consumption subject to the levy, by product.
* `fuel_policy_pages`: change monitor for the government pages that describe fuel supports (hash and metadata only, no page text).
"""
from __future__ import annotations

import html as _html
import io
import json
import re
import zipfile
from datetime import date, datetime

import httpx
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from s33weather.models import sha256

from ..db import utcnow
from .base import Collector, Result, get_with_retry, log_fetch
from .tables import Tally, csv_bytes, decode, gz, load_state, publish_table, save_state

BROWSER = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36"


def fetch(client: httpx.Client, url: str, timeout: float = 180) -> httpx.Response:
    """GET with the hub's own user agent; if a publisher's bot filter answers 403, one retry with an ordinary browser user agent (SEAI, Revenue)."""
    try:
        return get_with_retry(client, url, timeout=timeout)
    except httpx.HTTPStatusError as e:
        if e.response.status_code != 403:
            raise
    r = client.get(url, timeout=timeout, headers={"User-Agent": BROWSER, "Accept": "*/*", "Accept-Language": "en-IE,en;q=0.9"})
    r.raise_for_status()
    return r


def _text(raw: bytes) -> str:
    return re.sub(r"\s+", " ", _html.unescape(re.sub(r"<script.*?</script>|<style.*?</style>|<[^>]+>", " ", raw.decode("utf-8", "replace"), flags=re.S))).strip()


def _num(v):
    if v is None or v == "" or (isinstance(v, str) and v.strip().lower() in ("n/a", "na", "-", ":")):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _iso(v) -> str | None:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


# ----------------------------------------------------------------------------- EU Weekly Oil Bulletin

OB_PAGE = "https://energy.ec.europa.eu/data-and-analysis/weekly-oil-bulletin_en"
OB_HOST = "https://energy.ec.europa.eu"
OB_PRODUCTS = {"euro95": "petrol", "diesel": "diesel", "heating_oil": "heating_oil", "fuel_oil_1": "fuel_oil_low_sulphur", "fuel_oil_2": "fuel_oil_high_sulphur", "LPG": "lpg"}
OB_UNIT = {"fuel_oil_low_sulphur": "EUR per tonne", "fuel_oil_high_sulphur": "EUR per tonne"}
OB_COUNTRY = {"EU": "EU27", "EUR": "EURO_AREA", "GR": "EL", "UK": "UK"}
PRICE_COL = re.compile(r"^([A-Z]{2,3})_price_(with|wo)_tax_(euro95|diesel|heating_oil|fuel_oil_1|fuel_oil_2|LPG)$")


def find_bulletin_links(page_html: str) -> dict[str, str]:
    """The history workbook link on the Oil Bulletin page (the file name carries 'Prices_History')."""
    out = {}
    for href in re.findall(r'href="([^"]+)"', page_html):
        h = _html.unescape(href)
        low = h.lower()
        if "prices_history" in low and ".xls" in low:
            out["history"] = h if h.startswith("http") else OB_HOST + h
        elif "duties_and_taxes" in low and ".xls" in low:
            out["taxes"] = h if h.startswith("http") else OB_HOST + h
    return out


def parse_bulletin_prices(rows_with: list[tuple], rows_wo: list[tuple]) -> list[list]:
    """Long rows [date, country, product, unit, price_with_tax, price_without_tax] from the two price sheets (header row 0, units row 1, data from row 2)."""
    def index(rows):
        head = rows[0]
        cols = {}
        for i, h in enumerate(head):
            m = PRICE_COL.match(str(h or ""))
            if m:
                cols[i] = (OB_COUNTRY.get(m.group(1), m.group(1)), OB_PRODUCTS[m.group(3)])
        data = {}
        for r in rows[2:]:
            d = _iso(r[0])
            if not d:
                continue
            for i, (c, p) in cols.items():
                if i < len(r):
                    v = _num(r[i])
                    if v is not None:
                        data[(d, c, p)] = round(v, 4)
        return data
    w, wo = index(rows_with), index(rows_wo)
    out = []
    for k in sorted(set(w) | set(wo)):
        d, c, p = k
        out.append([d, c, p, OB_UNIT.get(p, "EUR per 1000 l"), w.get(k), wo.get(k)])
    return out


def parse_bulletin_country_taxes(rows: list[tuple], country_code: str = "IE_") -> list[list]:
    """One country's block from a tax sheet (VAT, Excise duties, Other Indirect Taxes): [since, petrol, diesel, heating_oil, fuel_oil_low_s, fuel_oil_high_s, lpg].
    Country code is on the first row of a block only; later rows of the block have none. A blank cell means the value did not change on that date."""
    out, cur = [], None
    for r in rows[4:]:
        if r and r[0]:
            cur = str(r[0])
        if cur != country_code or not r or len(r) < 3:
            continue
        d = _iso(r[1])
        if d:
            out.append([d] + [_num(v) for v in list(r[2:8]) + [None] * (6 - len(r[2:8]))])
    return out


class EuOilBulletin(Collector):
    key = "eu_oil_bulletin"
    name = "EU Weekly Oil Bulletin: consumer fuel prices and taxes, all member states, from 2005"
    publisher = "European Commission, Directorate-General for Energy"
    url = OB_PAGE
    licence = "Reproduction authorised provided the source is acknowledged (Oil Bulletin copyright notice, European Communities)"
    provides = ("Weekly pump prices with and without taxes for Euro-super 95, automotive diesel, heating gas oil, fuel oils and LPG in every member state, and Ireland's VAT, "
                "excise and other indirect tax tables as the Commission records them. Prices are euro per 1000 litres (per tonne for fuel oil).")
    used_by = ("fuelwatch",)
    interval_hours = 24.0
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        page = fetch(client, OB_PAGE, 90)
        links = find_bulletin_links(page.text)
        if "history" not in links:
            raise RuntimeError("the Oil Bulletin history workbook link was not found on the page")
        raw = fetch(client, links["history"], 300).content
        log_fetch(session, self.key, links["history"], sha256(raw), len(raw), self.licence, note="prices history workbook")
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        rows_with = list(wb["Prices with taxes"].iter_rows(values_only=True))
        rows_wo = list(wb["Prices wo taxes"].iter_rows(values_only=True))
        long = parse_bulletin_prices(rows_with, rows_wo)
        ie = [r for r in long if r[1] == "IE" and r[2] == "petrol" and r[4] is not None]
        countries = {r[1] for r in long}
        if len(ie) < 900 or len(countries) < 25:
            raise ValueError(f"Oil Bulletin looks wrong: {len(ie)} Irish petrol weeks, {len(countries)} countries")
        publish_table(items, "prices", "v1/fuel/oil_bulletin/prices.csv.gz",
                      gz(csv_bytes(["date", "country", "product", "unit", "price_with_tax", "price_without_tax"], long)), "application/gzip", len(long),
                      "EU Weekly Oil Bulletin prices with and without taxes, all member states, long format", self.key,
                      dict(source_url=links["history"], source_sha256=sha256(raw)))
        for sheet, name in (("VAT", "vat"), ("Excise duties", "excise"), ("Other Indirect Taxes", "other")):
            tab = parse_bulletin_country_taxes(list(wb[sheet].iter_rows(values_only=True)))
            if not tab:
                raise ValueError(f"no Irish rows in the {sheet} sheet")
            publish_table(items, "ie_" + name, f"v1/fuel/oil_bulletin/ie_{name}.csv.gz",
                          gz(csv_bytes(["since", "petrol", "diesel", "heating_oil", "fuel_oil_low_sulphur", "fuel_oil_high_sulphur", "lpg"], tab)), "application/gzip", len(tab),
                          f"Ireland {sheet} as recorded in the EU Oil Bulletin (euro per 1000 l, VAT in percent), by date of entry into force", self.key,
                          dict(source_url=links["history"], source_sha256=sha256(raw)))
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = len(long), f"{len(long):,} price rows, {len(countries)} areas, latest Irish petrol week {ie[-1][0]}"
        return res


# ----------------------------------------------------------------------------- Revenue Mineral Oil Tax and carbon taxes

MOT_PAGE = "https://www.revenue.ie/en/companies-and-charities/excise-and-licences/excise-duty-rates/mineral-oil-tax.aspx"
MOT_MANUAL = "https://www.revenue.ie/en/tax-professionals/tdm/excise/excise-duty-rates/energy-excise-duty-rates.pdf"
MONTHS = {m: i + 1 for i, m in enumerate("January February March April May June July August September October November December".split())}
HEAD = re.compile(r"Mineral\s+Oil\s+Tax\s+Rates\s+effective\s+from\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})(?:\s+to\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4}))?", re.I)


def _d(day: str, mon: str, year: str) -> str:
    return date(int(year), MONTHS[mon.capitalize()], int(day)).isoformat()


def _triple(line: str) -> list[float] | None:
    """The last three euro amounts on a row (non-carbon, carbon, total), or the single total of a pre-carbon-tax row."""
    tail = line.split("…")[-1] if "…" in line else line
    nums = [float(n) for n in re.findall(r"(\d{1,3}\.\d{2})\*?", tail)]
    if len(nums) >= 3:
        return nums[-3:]
    if len(nums) == 1:
        return [nums[0], 0.0, nums[0]]
    return None


def parse_mot(text: str) -> list[dict]:
    """Mineral Oil Tax per 1,000 litres for each period in Revenue's rates manual: petrol, auto-diesel (heavy oil used as a propellant), kerosene (used other than
    as a propellant) and marked gas oil (other heavy oil). Each row is checked: non-carbon + carbon = total."""
    heads = list(HEAD.finditer(text))
    out = []
    for i, h in enumerate(heads):
        seg = text[h.end(): heads[i + 1].start() if i + 1 < len(heads) else len(text)]
        if i == len(heads) - 1:
            seg = re.split(r"Solid\s+Fuel\s+Carbon\s+Tax\s+Rates", seg)[0]
        start = _d(h.group(1), h.group(2), h.group(3))
        end = _d(h.group(4), h.group(5), h.group(6)) if h.group(4) else None
        found: dict[str, list[float]] = {}
        for line in (re.sub(r"\s+", " ", l) for l in seg.splitlines()):
            if re.match(r"^\s*Petrol\b", line) and "petrol" not in found:
                t = _triple(line)
                if t:
                    found["petrol"] = t
            elif re.search(r"Used as a propellant", line) and "instead" not in line and "Liquefied" not in line and "diesel" not in found:
                t = _triple(line)
                if t:
                    found["diesel"] = t
            elif re.match(r"^\s*Kerosene", line) and "kerosene" not in found:
                t = _triple(line)
                if t:
                    found["kerosene"] = t
            elif re.match(r"^\s*Other heavy oil", line) and "marked_gas_oil" not in found:
                t = _triple(line)
                if t:
                    found["marked_gas_oil"] = t
        if "petrol" not in found or "diesel" not in found:
            continue
        for prod, (nc, c, tot) in found.items():
            if abs(nc + c - tot) > 0.011:
                raise ValueError(f"MOT {start} {prod}: {nc} + {c} does not equal {tot}")
            out.append(dict(valid_from=start, valid_to=end, product=prod, non_carbon=nc, carbon=c, total=tot))
    return out


def parse_carbon_schedule(text: str, heading: str, columns: list[str]) -> list[dict]:
    """A 'With effect from' table (rows: date, euro per tonne of CO2, then one value per column) following `heading`."""
    i = [m.start() for m in re.finditer(heading, text)]
    if not i:
        return []
    seg = text[i[-1]: i[-1] + 6000]
    rows = []
    for line in seg.splitlines():
        m = re.match(r"^\s*(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})\s+(.*)$", line)
        if not m or m.group(2).capitalize() not in MONTHS:
            continue
        vals = [float(v.replace(",", "")) for v in re.findall(r"€?\s*([\d,]+\.\d{2,4})", m.group(4))]
        if len(vals) == len(columns) + 1:
            rows.append(dict(valid_from=_d(*m.group(1, 2, 3)), per_tonne_co2=vals[0], **dict(zip(columns, vals[1:]))))
    return rows


def parse_mot_page(page_text: str) -> dict:
    """Current petrol and auto-diesel totals from the Revenue page text, for a cross-check against the manual."""
    t = re.sub(r"\s+", " ", page_text)
    eff = re.search(r"effective from (\d{1,2}) ([A-Za-z]+) (\d{4})", t)
    petrol = re.search(r"Petrol\s*(?:€|�)?\s*([\d.]+)\s*(?:€|�)?\s*([\d.]+)\s*(?:€|�)?\s*([\d.]+)", t)
    diesel = re.search(r"Used as a propellant\s*(?:€|�)?\s*([\d.]+)\s*(?:€|�)?\s*([\d.]+)\s*(?:€|�)?\s*([\d.]+)", t)
    pub = re.search(r"Published: (\d{1,2} [A-Za-z]+ \d{4})", t)
    return dict(effective=_d(*eff.groups()) if eff else None, published=pub.group(1) if pub else None,
                petrol=[float(x) for x in petrol.groups()] if petrol else None, diesel=[float(x) for x in diesel.groups()] if diesel else None)


class RevenueMot(Collector):
    key = "revenue_mot"
    name = "Revenue: Mineral Oil Tax rates since 2008 and the solid fuel and natural gas carbon tax schedules"
    publisher = "Revenue Commissioners"
    url = MOT_PAGE
    licence = "Published official rates (Revenue Tax and Duty Manual, Excise Duty Rates on Energy Products); check Revenue's terms for redistribution"
    provides = ("Mineral Oil Tax per 1,000 litres for petrol, auto-diesel, kerosene and marked gas oil for every rate period since 1 November 2008, split into non-carbon and carbon "
                "components, and the scheduled solid fuel and natural gas carbon tax rates. The hub checks that the parts add up to the total and that the manual agrees with Revenue's current page.")
    used_by = ("fuelwatch",)
    interval_hours = 24.0
    rights = "check"
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        from pypdf import PdfReader
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        pdf = fetch(client, MOT_MANUAL, 240).content
        log_fetch(session, self.key, MOT_MANUAL, sha256(pdf), len(pdf), self.licence, note="Excise Duty Rates manual")
        text = "\n".join(p.extract_text(extraction_mode="layout") or "" for p in PdfReader(io.BytesIO(pdf)).pages)
        rows = parse_mot(text)
        periods = sorted({r["valid_from"] for r in rows})
        if len(periods) < 25 or periods[-1] < "2026-01-01":
            raise ValueError(f"the Revenue manual parsed to only {len(periods)} periods (latest {periods[-1] if periods else None})")
        page = fetch(client, MOT_PAGE, 90).content
        log_fetch(session, self.key, MOT_PAGE, sha256(page), len(page), self.licence, note="current rates page")
        cur = parse_mot_page(_text(page))
        latest = {r["product"]: r for r in rows if r["valid_from"] == periods[-1]}
        agrees = bool(cur["petrol"] and cur["diesel"] and abs(cur["petrol"][2] - latest["petrol"]["total"]) < 0.011 and abs(cur["diesel"][2] - latest["diesel"]["total"]) < 0.011)
        if cur["petrol"] and not agrees:
            raise ValueError(f"Revenue's page ({cur['petrol']}, {cur['diesel']}) does not match the manual ({latest['petrol']}, {latest['diesel']}); not publishing")
        publish_table(items, "mot", "v1/fuel/revenue/mot_rates.csv.gz",
                      gz(csv_bytes(["valid_from", "valid_to", "product", "non_carbon", "carbon", "total"], [[r["valid_from"], r["valid_to"] or "", r["product"], r["non_carbon"], r["carbon"], r["total"]] for r in rows])),
                      "application/gzip", len(rows), "Mineral Oil Tax per 1,000 litres by rate period (euro), non-carbon and carbon components", self.key,
                      dict(source_url=MOT_MANUAL, source_sha256=sha256(pdf)))
        sfct = parse_carbon_schedule(text, r"Amount\s+SFCT\s+rate", ["coal", "peat_briquettes", "milled_peat", "other_peat"])
        ngct = parse_carbon_schedule(text, r"Natural\s+Gas\s+Carbon\s+Tax\s+Rate", ["rate_per_mwh_gcv", "ncv_to_gcv_factor"])
        publish_table(items, "sfct", "v1/fuel/revenue/solid_fuel_carbon_tax.csv.gz", gz(csv_bytes(["valid_from", "per_tonne_co2", "coal", "peat_briquettes", "milled_peat", "other_peat"],
                      [[r["valid_from"], r["per_tonne_co2"], r["coal"], r["peat_briquettes"], r["milled_peat"], r["other_peat"]] for r in sfct])), "application/gzip", len(sfct),
                      "Solid Fuel Carbon Tax per tonne of fuel, current and scheduled (euro)", self.key, dict(source_url=MOT_MANUAL, source_sha256=sha256(pdf)))
        publish_table(items, "ngct", "v1/fuel/revenue/natural_gas_carbon_tax.csv.gz", gz(csv_bytes(["valid_from", "per_tonne_co2", "rate_per_mwh_gcv", "ncv_to_gcv_factor"],
                      [[r["valid_from"], r["per_tonne_co2"], r["rate_per_mwh_gcv"], r["ncv_to_gcv_factor"]] for r in ngct])), "application/gzip", len(ngct),
                      "Natural Gas Carbon Tax per MWh at gross calorific value, current and scheduled (euro)", self.key, dict(source_url=MOT_MANUAL, source_sha256=sha256(pdf)))
        if len(sfct) < 8 or len(ngct) < 8:
            raise ValueError(f"carbon tax schedules parsed to {len(sfct)} and {len(ngct)} rows")
        meta = dict(manual_url=MOT_MANUAL, manual_sha256=sha256(pdf), page_url=MOT_PAGE, page_sha256=sha256(page), page_published=cur["published"], page_effective=cur["effective"],
                    page_agrees_with_manual=agrees, periods=len(periods), first_period=periods[0], latest_period=periods[-1], sfct_rows=len(sfct), ngct_rows=len(ngct),
                    retrieved_on=utcnow().date().isoformat(), note="Parsed from the manual's tables; each row is checked non-carbon + carbon = total. The manual and Revenue's page are the authority.")
        publish_table(items, "meta", "v1/fuel/revenue/mot_meta.json", json.dumps(meta, indent=1, sort_keys=True).encode(), "application/json", len(periods), "How the Revenue rates were read and checked", self.key)
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = len(rows), f"{len(periods)} rate periods ({periods[0]} to {periods[-1]}), page agrees with manual: {agrees}"
        return res


# ----------------------------------------------------------------------------- EIA Brent and ECB reference rates

EIA_URL = "https://www.eia.gov/dnav/pet/hist_xls/RBRTEd.xls"
ECB_URL = "https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip"


def parse_brent(rows: list[list]) -> list[list]:
    """[date, usd_per_barrel] from the EIA 'Data 1' sheet rows (xlrd values with dates already converted to ISO strings)."""
    out = []
    for d, v in rows:
        if d and v not in ("", None):
            out.append([d, round(float(v), 4)])
    return out


class EiaBrent(Collector):
    key = "eia_brent"
    name = "US EIA Europe Brent spot price FOB, daily"
    publisher = "U.S. Energy Information Administration"
    url = "https://www.eia.gov/dnav/pet/hist/RBRTED.htm"
    licence = "EIA data are public information; cite the U.S. Energy Information Administration as the source"
    provides = "Daily Europe Brent spot price in US dollars per barrel from 1987."
    used_by = ("fuelwatch",)
    interval_hours = 24.0
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        import xlrd
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        raw = fetch(client, EIA_URL, 120).content
        log_fetch(session, self.key, EIA_URL, sha256(raw), len(raw), self.licence)
        wb = xlrd.open_workbook(file_contents=raw)
        ws = wb.sheet_by_name("Data 1")
        pairs = []
        for r in range(3, ws.nrows):
            d, v = ws.cell_value(r, 0), ws.cell_value(r, 1)
            if isinstance(d, float) and v not in ("", None):
                y, m, dd = xlrd.xldate_as_tuple(d, wb.datemode)[:3]
                pairs.append([date(y, m, dd).isoformat(), v])
        rows = parse_brent(pairs)
        if len(rows) < 8000 or rows[-1][0] < "2026-01-01":
            raise ValueError(f"Brent file looks wrong: {len(rows)} rows, last {rows[-1][0] if rows else None}")
        publish_table(items, "brent", "v1/fuel/eia/brent_daily.csv.gz", gz(csv_bytes(["date", "usd_per_barrel"], rows)), "application/gzip", len(rows),
                      "Europe Brent spot price FOB, US dollars per barrel, daily (EIA)", self.key, dict(source_url=EIA_URL, source_sha256=sha256(raw)))
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = len(rows), f"{len(rows):,} days, latest {rows[-1][0]} at ${rows[-1][1]}"
        return res


def parse_ecb(csv_text: str, columns: tuple[str, ...] = ("USD", "GBP")) -> list[list]:
    lines = csv_text.strip().splitlines()
    head = [h.strip() for h in lines[0].split(",")]
    idx = {c: head.index(c) for c in columns}
    out = []
    for ln in lines[1:]:
        p = [x.strip() for x in ln.split(",")]
        row = [p[0]] + [(_num(p[idx[c]]) if p[idx[c]] not in ("N/A", "") else None) for c in columns]
        if any(v is not None for v in row[1:]):
            out.append(row)
    return sorted(out)


class EcbFx(Collector):
    key = "ecb_fx"
    name = "ECB euro foreign exchange reference rates (US dollar, sterling)"
    publisher = "European Central Bank"
    url = "https://www.ecb.europa.eu/stats/policy_and_exchange_rates/euro_reference_exchange_rates/html/index.en.html"
    licence = "ECB statistics: reproduction permitted provided the source is acknowledged"
    provides = "Daily euro reference rates against the US dollar and sterling from 1999 (units of currency per euro)."
    used_by = ("fuelwatch",)
    interval_hours = 24.0
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        raw = fetch(client, ECB_URL, 120).content
        log_fetch(session, self.key, ECB_URL, sha256(raw), len(raw), self.licence)
        z = zipfile.ZipFile(io.BytesIO(raw))
        rows = parse_ecb(z.read(z.namelist()[0]).decode("utf-8", "replace"))
        if len(rows) < 6000 or rows[-1][0] < "2026-01-01":
            raise ValueError(f"ECB file looks wrong: {len(rows)} rows, last {rows[-1][0] if rows else None}")
        publish_table(items, "fx", "v1/fuel/ecb/eur_usd_gbp.csv.gz", gz(csv_bytes(["date", "usd_per_eur", "gbp_per_eur"], rows)), "application/gzip", len(rows),
                      "ECB euro reference rates: US dollar and sterling per euro, daily", self.key, dict(source_url=ECB_URL, source_sha256=sha256(raw)))
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = len(rows), f"{len(rows):,} days, latest {rows[-1][0]} USD {rows[-1][1]}"
        return res


# ----------------------------------------------------------------------------- SEAI

SEAI = "https://www.seai.ie"
SEAI_FILES = {
    "archive": "/sites/default/files/publications/Domestic-Fuel-Cost-Archive.xlsx",
    "final": "/sites/default/files/publications/Final-Energy-Consumption.xlsx",
    "by_fuel": "/sites/default/files/publications/Energy-by-Fuel.xlsx",
    "balance": "/sites/default/files/2025-05/National-Energy-Balance.xlsx",
    "factors": "/sites/default/files/data-and-insights/seai-statistics/conversion-factors/SEAI-conversion-and-emission-factors.xlsx",
}
SEAI_DOWNLOADS = SEAI + "/data-and-insights/seai-statistics/energy-data-downloads"


def seai_links(page_html: str) -> dict[str, str]:
    """Current download links from the SEAI Energy Data Downloads page (the National Energy Balance path carries a year-month)."""
    out = {}
    for href in re.findall(r'href="([^"]+\.xlsx)"', page_html):
        low = href.lower()
        for key, tag in (("archive", "domestic-fuel-cost-archive"), ("final", "final-energy-consumption"), ("by_fuel", "energy-by-fuel"), ("balance", "national-energy-balance")):
            if tag in low:
                out[key] = href if href.startswith("http") else SEAI + href
    return out


def parse_archive(rows_unit: list[tuple], rows_ckwh: list[tuple]) -> list[list]:
    """Long rows [date, fuel, euro_per_unit, cent_per_kwh] from the two archive sheets (header row 0: Date then one column per fuel; text values such as 'n/a' are dropped)."""
    def grid(rows):
        head = [str(h).strip() if h else "" for h in rows[0]]
        out = {}
        for r in rows[1:]:
            d = _iso(r[0])
            if not d:
                continue
            for i, h in enumerate(head[1:], 1):
                if h and h != "Special Notes" and i < len(r) and _num(r[i]) is not None:
                    out[(d, re.sub(r"\s+", " ", h))] = _num(r[i])
        return out
    a, b = grid(rows_unit), grid(rows_ckwh)
    return [[d, f, round(a[(d, f)], 5) if (d, f) in a else None, round(b[(d, f)], 5) if (d, f) in b else None] for d, f in sorted(set(a) | set(b))]


def parse_year_table(rows: list[tuple], sheet: str) -> list[list]:
    """SEAI year-by-row tables (header: label, NACE, 1990, 1991, ...): [sheet, label, nace, year, value]."""
    head = rows[0]
    years = {i: int(h) for i, h in enumerate(head) if isinstance(h, (int, float)) and 1980 < h < 2100}
    out = []
    for r in rows[1:]:
        label = (str(r[0]).strip() if r and r[0] is not None else "")
        if not label:
            continue
        for i, y in years.items():
            if i < len(r) and _num(r[i]) is not None:
                out.append([sheet, re.sub(r"\s+", " ", label), (str(r[1]).strip() if len(r) > 1 and r[1] is not None else ""), y, round(_num(r[i]), 4)])
    return out


def parse_balance_sheet(rows: list[tuple], sheet: str) -> list[list]:
    """One National Energy Balance sheet (header: title, NACE, fuel columns; rows: flows): [year, status, flow, fuel, ktoe]."""
    m = re.match(r"\s*(\d{4})\s*(.*?)\s*(?:Units|$)", str(rows[0][0]) if rows and rows[0] and rows[0][0] else sheet, re.S)
    year = int(m.group(1)) if m else int(re.match(r"\d{4}", sheet).group(0))
    status = "interim" if "interim" in sheet.lower() else "final"
    head = [re.sub(r"\s+", " ", str(h)).strip() if h else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        flow = re.sub(r"\s+", " ", str(r[0])).strip() if r and r[0] else ""
        if not flow:
            continue
        for i in range(2, min(len(head), len(r))):
            if head[i] and _num(r[i]) is not None:
                out.append([year, status, flow, head[i], round(_num(r[i]), 4)])
    return out


class SeaiFuel(Collector):
    key = "seai_fuel"
    name = "SEAI: domestic fuel price archive, energy by fuel, final energy consumption and National Energy Balance"
    publisher = "Sustainable Energy Authority of Ireland (SEAI)"
    url = SEAI_DOWNLOADS
    licence = "SEAI statistics (Creative Commons Attribution 4.0 for SEAI open data; check the page for each download); attribute SEAI"
    provides = ("Quarterly delivered prices for coal products, oils, LPG, wood fuels, gas and electricity (euro per unit and cent per kWh) back to 1990, annual final energy consumption "
                "by sector and fuel, energy by fuel, the National Energy Balance by year, and the SEAI conversion factors. Converted from Excel to CSV.")
    used_by = ("fuelwatch",)
    interval_hours = 168.0
    rights = "check"
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        tally = Tally()
        links = {k: SEAI + p for k, p in SEAI_FILES.items()}
        try:
            links.update(seai_links(fetch(client, SEAI_DOWNLOADS, 90).text))     # the page's current links win over the built-in paths
        except httpx.HTTPError:
            pass

        def job(name, fn):
            try:
                fn()
            except Exception as e:  # noqa: BLE001 - one workbook failing must not stop the rest
                tally.failed.append(f"{name} ({type(e).__name__}: {str(e)[:80]})")

        def book(key):
            raw = fetch(client, links[key], 240).content
            log_fetch(session, self.key, links[key], sha256(raw), len(raw), self.licence, note=key)
            return load_workbook(io.BytesIO(raw), read_only=True, data_only=True), raw

        def pub(name, path, head, rows, desc, raw, url):
            if publish_table(items, name, path, gz(csv_bytes(head, rows)), "application/gzip", len(rows), desc, self.key, dict(source_url=url, source_sha256=sha256(raw))):
                tally.published += 1
            else:
                tally.unchanged += 1

        def archive():
            wb, raw = book("archive")
            rows = parse_archive(list(wb["Archive-D"].iter_rows(values_only=True)), list(wb["Archive-D ckWh"].iter_rows(values_only=True)))
            if len(rows) < 2000 or max(r[0] for r in rows) < "2026-01-01":
                raise ValueError(f"archive has {len(rows)} rows, latest {max(r[0] for r in rows) if rows else None}")
            pub("archive", "v1/fuel/seai/domestic_fuel_archive.csv.gz", ["date", "fuel", "euro_per_unit", "cent_per_kwh"], rows,
                "SEAI domestic fuel prices by quarter: euro per unit and cent per kWh delivered, VAT inclusive", raw, links["archive"])

        def final():
            wb, raw = book("final")
            rows = [r for ws in wb.worksheets for r in parse_year_table(list(ws.iter_rows(values_only=True)), ws.title)]
            if len(rows) < 5000:
                raise ValueError(f"final energy consumption has {len(rows)} rows")
            pub("final", "v1/fuel/seai/final_energy_consumption.csv.gz", ["sector", "fuel", "nace", "year", "ktoe"], rows, "SEAI final energy consumption by sector and fuel, ktoe", raw, links["final"])

        def by_fuel():
            wb, raw = book("by_fuel")
            rows = [r for ws in wb.worksheets for r in parse_year_table(list(ws.iter_rows(values_only=True)), ws.title)]
            if len(rows) < 3000:
                raise ValueError(f"energy by fuel has {len(rows)} rows")
            pub("by_fuel", "v1/fuel/seai/energy_by_fuel.csv.gz", ["fuel_group", "flow", "nace", "year", "ktoe"], rows, "SEAI energy balance flows by fuel group and year, ktoe", raw, links["by_fuel"])

        def balance():
            wb, raw = book("balance")
            rows = [r for ws in wb.worksheets if re.match(r"\d{4}", ws.title) for r in parse_balance_sheet(list(ws.iter_rows(values_only=True)), ws.title)]
            if len(rows) < 20000:
                raise ValueError(f"National Energy Balance has {len(rows)} rows")
            pub("balance", "v1/fuel/seai/national_energy_balance.csv.gz", ["year", "status", "flow", "fuel", "ktoe"], rows, "SEAI National Energy Balance by year, flow and fuel, ktoe", raw, links["balance"])

        def factors():
            wb, raw = book("factors")
            ws = wb["Conversion and emission factors"]
            rows = [[("" if c is None else (round(c, 6) if isinstance(c, float) else str(c).strip())) for c in r[:14]] for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r[:14])]
            pub("factors", "v1/fuel/seai/conversion_factors.csv.gz", [f"c{i}" for i in range(14)], rows, "SEAI conversion and emission factors sheet, as published", raw, links["factors"])

        for n, f in (("archive", archive), ("final", final), ("by_fuel", by_fuel), ("balance", balance), ("factors", factors)):
            job(n, f)
        if tally.published + tally.unchanged == 0:
            raise RuntimeError("no SEAI file could be read: " + "; ".join(tally.failed))
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = tally.published + tally.unchanged, tally.message(5)
        return res


# ----------------------------------------------------------------------------- NORA

NORA_PAGE = "https://www.nora.ie/volumes-of-oil-consumption"


def nora_xls_link(page_html: str) -> str | None:
    for href in re.findall(r'href="([^"]+\.xls[^"]*)"', page_html):
        h = _html.unescape(href)
        if "volumes" in h.lower() or "nora" in h.lower() or "_files/ugd" in h:
            return h.split("?")[0] if "?" in h and "dn=" in h else h
    return None


def parse_nora_sheet(rows: list[list], year: int) -> list[list]:
    """Monthly litres by product from one year's sheet. Two header rows (a qualifier such as 'BIOFUEL IN' above the product); month rows follow."""
    hdr_i = next(i for i, r in enumerate(rows) if len(r) > 1 and str(r[1]).strip().lower() == "month")
    names = []
    for j in range(len(rows[hdr_i])):
        top = str(rows[hdr_i - 1][j]).strip() if hdr_i > 0 and j < len(rows[hdr_i - 1]) and rows[hdr_i - 1][j] not in ("", None) else ""
        names.append((top + " " + str(rows[hdr_i][j]).strip()).strip() if str(rows[hdr_i][j]).strip() else "")
    out = []
    for r in rows[hdr_i + 1:]:
        m = str(r[1]).strip() if len(r) > 1 else ""
        if m not in MONTHS:
            if m.lower() == "total":
                break
            continue
        vals = [(names[j], _num(r[j])) for j in range(2, min(len(names), len(r))) if names[j] and _num(r[j]) is not None]
        if not any(v for _, v in vals):
            continue                                       # a month not yet reported: blank or zero in every column
        out.extend([f"{year}-{MONTHS[m]:02d}", re.sub(r"\s+", " ", n), round(v, 1)] for n, v in vals)
    return out


class NoraVolumes(Collector):
    key = "nora_volumes"
    name = "NORA: monthly volumes of oil consumption subject to the levy"
    publisher = "National Oil Reserves Agency (NORA)"
    url = NORA_PAGE
    licence = "Published by NORA; check the NORA terms of use before redistribution"
    provides = "Litres of gasoline, kerosene, gas oil, motor diesel and fuel oil released for consumption each month since 2008, including the biofuel share where reported."
    used_by = ("fuelwatch",)
    interval_hours = 168.0
    rights = "check"
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        import xlrd
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        page = fetch(client, NORA_PAGE, 90)
        link = nora_xls_link(page.text)
        if not link:
            raise RuntimeError("the NORA volumes workbook link was not found on the page")
        raw = fetch(client, link, 240).content
        log_fetch(session, self.key, link, sha256(raw), len(raw), self.licence)
        wb = xlrd.open_workbook(file_contents=raw)
        rows = []
        for s in wb.sheets():
            m = re.search(r"(\d{2,4})\s*$", s.name)
            if not m:
                continue
            y = int(m.group(1)); y = y + 2000 if y < 100 else y
            rows += parse_nora_sheet([s.row_values(i) for i in range(s.nrows)], y)
        months = sorted({r[0] for r in rows})
        if len(months) < 150 or months[-1] < "2026-01":
            raise ValueError(f"NORA workbook parsed to {len(months)} months, latest {months[-1] if months else None}")
        publish_table(items, "volumes", "v1/fuel/nora/volumes_monthly.csv.gz", gz(csv_bytes(["month", "product", "litres"], rows)), "application/gzip", len(rows),
                      "NORA monthly volumes of oil consumption subject to the levy, litres by product", self.key, dict(source_url=link, source_sha256=sha256(raw)))
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = len(rows), f"{len(months)} months ({months[0]} to {months[-1]})"
        return res


# ----------------------------------------------------------------------------- policy page monitor

POLICY_PAGES = {
    "dot_support": ("Department of Transport: What the Government is doing", "https://www.gov.ie/en/department-of-transport/campaigns/what-the-government-is-doing/"),
    "dsp_fuel_allowance": ("Department of Social Protection: Fuel Allowance", "https://www.gov.ie/en/department-of-social-protection/services/fuel-allowance/"),
    "revenue_mot_page": ("Revenue: Mineral Oil Tax", MOT_PAGE),
    "revenue_energy_taxes": ("Revenue: Excise duty rates on natural gas, solid fuel and electricity", "https://www.revenue.ie/en/companies-and-charities/excise-and-licences/excise-duty-rates/index.aspx"),
}


def page_signature(raw: bytes) -> dict:
    """What the monitor keeps: the title, a hash of the visible text and a 'last updated' or 'published' date if the page states one. No page text."""
    s = raw.decode("utf-8", "replace")
    title = re.search(r"<title>(.*?)</title>", s, re.S)
    text = _text(raw)
    upd = re.search(r"(?:Last updated|Updated|Published)[: ]+(\d{1,2} [A-Za-z]+ \d{4})", text)
    return dict(title=re.sub(r"\s+", " ", _html.unescape(title.group(1))).strip() if title else "", text_sha256=sha256(text.encode()), text_chars=len(text), stated_date=upd.group(1) if upd else None)


class FuelPolicyPages(Collector):
    key = "fuel_policy_pages"
    name = "Fuel support and tax pages: change monitor"
    publisher = "Government of Ireland (gov.ie), Revenue"
    url = "https://www.gov.ie/"
    licence = "Metadata only: page title, a hash of the visible text, and the date the page states. No page text is kept."
    provides = "Whether the government pages that describe fuel supports and fuel taxes have changed since an editor last reviewed them."
    used_by = ("fuelwatch",)
    interval_hours = 24.0
    rights = "check"
    domain, tier = "fuel", "files"

    def run(self, session: Session, client: httpx.Client) -> Result:
        state = load_state(session, self.key)
        items = state.setdefault("items", {})
        out, failed = [], []
        for key, (title, url) in POLICY_PAGES.items():
            try:
                raw = fetch(client, url, 90).content
                log_fetch(session, self.key, url, sha256(raw), len(raw), self.licence, note=key)
                sig = page_signature(raw)
                prev = state.get("pages", {}).get(key, {})
                changed_at = utcnow().isoformat(timespec="seconds") if prev.get("text_sha256") != sig["text_sha256"] else prev.get("changed_at")
                state.setdefault("pages", {})[key] = dict(sig, changed_at=changed_at)
                out.append(dict(key=key, name=title, url=url, status="ok", checked_at=utcnow().isoformat(timespec="seconds"), changed_at=changed_at, **sig))
            except Exception as e:  # noqa: BLE001 - recorded as unavailable, never as unchanged
                prev = state.get("pages", {}).get(key, {})
                failed.append(f"{key} ({type(e).__name__})")
                out.append(dict(key=key, name=title, url=url, status="unavailable", checked_at=utcnow().isoformat(timespec="seconds"), changed_at=prev.get("changed_at"),
                                title=prev.get("title", ""), text_sha256=prev.get("text_sha256"), text_chars=prev.get("text_chars"), stated_date=prev.get("stated_date")))
        publish_table(items, "pages", "v1/fuel/policy/page_monitor.json", json.dumps(dict(generated_at=utcnow().isoformat(timespec="seconds"), pages=out), indent=1, sort_keys=True).encode(),
                      "application/json", len(out), "Change monitor for fuel support and tax pages (title, text hash, stated date)", self.key)
        res = save_state(session, self.key, state, self.url, self.name)
        res.fetched, res.message = len(out) - len(failed), f"{len(out) - len(failed)} of {len(out)} pages read" + (f"; unavailable: {', '.join(failed)}" if failed else "")
        if len(failed) == len(out):
            raise RuntimeError("no policy page could be read: " + "; ".join(failed))
        return res
